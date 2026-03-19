"""Extract text content from .docx files."""

from io import BytesIO

from docx import Document


def extract_docx(content_bytes: bytes, include_tables: bool = True) -> str:
    """Extract text from a .docx file, preserving document order.

    Args:
        content_bytes: Raw bytes of the .docx file.
        include_tables: If True, render tables as markdown tables.

    Returns:
        Extracted text content.
    """
    doc = Document(BytesIO(content_bytes))
    parts = []

    for element in doc.element.body:
        if element.tag.endswith("}p"):
            para = _find_paragraph(doc, element)
            if para and para.text.strip():
                parts.append(para.text)
        elif element.tag.endswith("}tbl") and include_tables:
            table = _find_table(doc, element)
            if table:
                md = _table_to_markdown(table)
                if md:
                    parts.append(md)

    return "\n\n".join(parts)


def _find_paragraph(doc: Document, element) -> object:
    """Find the python-docx Paragraph object for an XML element."""
    for para in doc.paragraphs:
        if para._element is element:
            return para
    return None


def _find_table(doc: Document, element) -> object:
    """Find the python-docx Table object for an XML element."""
    for table in doc.tables:
        if table._element is element:
            return table
    return None


def _table_to_markdown(table) -> str:
    """Convert a python-docx Table to a markdown table string."""
    rows = []
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        rows.append("| " + " | ".join(cells) + " |")

    if len(rows) >= 1:
        col_count = len(table.rows[0].cells)
        separator = "| " + " | ".join(["---"] * col_count) + " |"
        rows.insert(1, separator)

    return "\n".join(rows)
